from pathlib import Path
from typing import Iterable
import io
import re

import fitz
from bs4 import BeautifulSoup
from docx import Document

from .config import ALLOWED_EXTENSIONS


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_file(path: Path) -> list[dict]:
    """Return [{page, text}]. Page is 1-based for PDFs; null for other formats."""
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")
    if suffix == ".pdf":
        pages = []
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf, start=1):
                text = clean_text(page.get_text("text"))
                if text:
                    pages.append({"page": i, "text": text})
        return pages
    if suffix == ".txt" or suffix == ".md":
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix in {".html", ".htm"}:
        text = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser").get_text("\n")
    elif suffix == ".docx":
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    text = clean_text(text)
    return [{"page": None, "text": text}] if text else []


def chunk_pages(pages: list[dict], max_chars: int = 1500, overlap: int = 200) -> list[dict]:
    """Split text into overlapping chunks without losing PDF page provenance."""
    chunks = []
    for page_item in pages:
        text = page_item["text"]
        start = 0
        while start < len(text):
            end = min(len(text), start + max_chars)
            if end < len(text):
                # Prefer ending at a sentence/paragraph boundary where possible.
                boundary = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
                if boundary > start + max_chars // 2:
                    end = boundary + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append({"page": page_item.get("page"), "text": piece})
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks


def extract_bytes(filename: str, data: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    # Parse in-memory upload through a temporary local file; removed immediately.
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as tmp:
        tmp.write(data)
        tmp.flush()
        return extract_file(Path(tmp.name))
